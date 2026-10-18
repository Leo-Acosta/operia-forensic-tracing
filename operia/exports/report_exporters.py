from pathlib import Path
import json
import re
from html import escape
from textwrap import wrap

import markdown as md
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from operia.reports.models import Presentation, Report
from operia.legal_report.services.text_service import normalize_spanish
from operia.legal_report.services.report_service import (
    NO_ATTRIBUTION_TEXT,
    attribution_assessment,
    build_attribution_conclusion,
    build_matrix_based_conclusion,
    build_smart_executive_summary,
    exportability_markdown,
    isolated_case_snapshot,
    legal_probatory_assessment,
    report_annexes_payload,
    report_exportability_assessment,
)
from operia.core.config import settings
from operia.shared.time import utcnow

from operia.exports.paths import new_export_path

PUBLIC_FORBIDDEN_MARKERS = (
    "ollama",
    "mistral",
    "fallback",
    " ia local",
    "asistente ia",
    "documento asistido por ia",
    "herramientas automatizadas",
    "la ia no debe",
    "no hay encuadre legal",
    "sin encuadre legal cargado",
    "enriquecido_con_ia",
    "```json",
)


def export_markdown(report: Report) -> Path:
    path = new_export_path(f"report-{report.id}.md")
    if report.template_used == "CPython AST revisado":
        path.write_bytes(report.markdown_content.encode("utf-8"))
    else:
        path.write_text(report.markdown_content if report.template_used == "Motor propio OPERIA" else normalize_spanish(report.markdown_content), encoding="utf-8")
    return path


def export_html(report: Report) -> Path:
    if report.template_used in ("Motor propio OPERIA", "CPython AST revisado"):
        path = new_export_path(f"report-{report.id}.html")
        path.write_text("<!doctype html><meta charset='utf-8'><pre>" + escape(report.markdown_content) + "</pre>", encoding="utf-8")
        return path
    html = "<pre>" + escape(normalize_spanish(report.markdown_content)) + "</pre>"
    path = new_export_path(f"report-{report.id}.html")
    path.write_text(f"<!doctype html><meta charset='utf-8'><body>{html}</body>", encoding="utf-8")
    return path


def _markdown_sections(markdown: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    current_title = "Informe"
    current_lines: list[str] = []
    for line in normalize_spanish(markdown).splitlines():
        if line.startswith("## "):
            if current_lines:
                sections.append((current_title, "\n".join(current_lines).strip()))
            current_title = line[3:].strip()
            current_lines = []
        elif line.startswith("# "):
            current_title = line[2:].strip()
        else:
            current_lines.append(line)
    if current_lines:
        sections.append((current_title, "\n".join(current_lines).strip()))
    return sections


def _clean_markdown_text(text: str) -> str:
    return (
        _public_text(text)
        .replace("**", "")
        .replace("`", "")
        .replace("| --- | --- | --- | --- | --- |", "")
        .strip()
    )


def export_docx(report: Report) -> Path:
    if report.template_used in ("Motor propio OPERIA", "CPython AST revisado"):
        document = Document()
        document.add_heading(report.title, 0)
        for line in report.markdown_content.splitlines():
            document.add_paragraph(line)
        path = new_export_path(f"report-{report.id}.docx")
        document.save(path)
        return path
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(15, 23, 42)
    styles["Heading 2"].font.color.rgb = RGBColor(21, 94, 117)

    logo = _logo_path()
    if logo:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(logo), width=Inches(3.6))

    title = document.add_heading(normalize_spanish(report.title), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Caso {report.case.case_code} | {normalize_spanish(report.case.incident_type)}").bold = True
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Evidencias", "Hallazgos", "Daño estimado", "Estado"]
    values = [
        str(len(report.case.evidence)),
        str(len(report.case.findings)),
        f"{sum(item.amount for item in report.case.damage_items):,.2f} ARS",
        normalize_spanish(report.status),
    ]
    for index, cell in enumerate(table.rows[0].cells):
        cell.text = f"{headers[index]}\n{values[index]}"

    document.add_page_break()
    for line in normalize_spanish(report.markdown_content).splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("|") and "---" not in line:
            document.add_paragraph(_clean_markdown_text(line))
        elif line.strip():
            document.add_paragraph(_clean_markdown_text(line))
    path = new_export_path(f"report-{report.id}.docx")
    document.save(path)
    return path


def export_pdf(report: Report) -> Path:
    path = new_export_path(f"report-{report.id}.pdf")
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    case = report.case
    logo = _logo_path()
    sections = _markdown_sections(report.markdown_content)
    total_damage = sum(item.amount for item in case.damage_items)

    def report_background(page_label: str = "", show_header_logo: bool = True) -> None:
        c.setFillColor(colors.white)
        c.rect(0, 0, width, height, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#f6f9fd"))
        c.roundRect(28, 28, width - 56, height - 56, 16, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.roundRect(44, 44, width - 88, height - 88, 14, fill=1, stroke=0)
        c.setStrokeColor(colors.HexColor("#dbe5f1"))
        c.roundRect(44, 44, width - 88, height - 88, 14, fill=0, stroke=1)
        c.setStrokeColor(colors.HexColor("#18a999"))
        c.setLineWidth(2)
        c.line(64, height - 78, width - 64, height - 78)
        if logo and show_header_logo:
            c.drawImage(ImageReader(str(logo)), 60, height - 68, width=120, height=44, preserveAspectRatio=True, mask="auto")
        c.setFont("Helvetica-Bold", 8)
        c.setFillColor(colors.HexColor("#64748b"))
        c.drawRightString(width - 60, height - 49, "INFORME FORENSE-JURÍDICO")
        if page_label:
            c.drawRightString(width - 60, 34, page_label)

    report_background(show_header_logo=False)
    if logo:
        c.drawImage(ImageReader(str(logo)), 72, height - 188, width=230, height=88, preserveAspectRatio=True, mask="auto")
    _draw_chip(c, 72, height - 210, "INFORME INTEGRAL", "#eef8f7", "#b9ebe4", "#087f5b")
    _draw_chip(c, 190, height - 210, "CONTROL PROFESIONAL", "#fff7ed", "#fed7aa", "#9a3412")
    c.setFont("Helvetica-Bold", 26)
    c.setFillColor(colors.HexColor("#0f172a"))
    _draw_wrapped(c, normalize_spanish(report.title), 72, height - 260, width - 144, "Helvetica-Bold", 26, 31, colors.HexColor("#0f172a"), max_lines=3)
    c.setFont("Helvetica", 12)
    c.setFillColor(colors.HexColor("#475467"))
    c.drawString(74, height - 330, normalize_spanish(f"Caso {case.case_code} | {case.incident_type}"))
    c.drawString(74, height - 350, normalize_spanish(f"Responsable: {_responsible_text(report)}"))
    c.drawString(74, height - 370, normalize_spanish(f"Jurisdicción: {case.jurisdiction}"))

    card_y = 142
    _draw_metric(c, 72, card_y, "Evidencias", str(len(case.evidence)), "#155e75")
    _draw_metric(c, 230, card_y, "Hallazgos", str(len(case.findings)), "#087f5b")
    _draw_metric(c, 388, card_y, "Daño estimado", f"{total_damage:,.0f} ARS", "#b7791f")
    c.setFont("Helvetica", 8)
    c.setFillColor(colors.HexColor("#64748b"))
    c.drawRightString(width - 72, 74, "OPERIA Forensic Tracing | Informe forense profesional")
    c.showPage()

    page = 1
    for title, content in sections:
        if not content and title == "Operia Forensic Legal Report":
            continue
        report_background(f"Página {page:02d}")
        c.setFont("Helvetica-Bold", 19)
        c.setFillColor(colors.HexColor("#0f172a"))
        y = height - 112
        y = _draw_wrapped(c, normalize_spanish(title), 68, y, width - 136, "Helvetica-Bold", 19, 23, colors.HexColor("#0f172a"), max_lines=2) - 10

        c.setFillColor(colors.HexColor("#f8fbff"))
        c.setStrokeColor(colors.HexColor("#dbe5f1"))
        c.roundRect(64, 76, width - 128, y - 86, 12, fill=1, stroke=1)
        c.setStrokeColor(colors.HexColor("#18a999"))
        c.setLineWidth(3)
        c.line(82, y - 22, 82, 98)
        y -= 25

        for paragraph in _clean_markdown_text(content).splitlines():
            paragraph = paragraph.strip()
            if not paragraph or paragraph.startswith("| ---"):
                y -= 6
                continue
            if y < 110:
                c.showPage()
                page += 1
                report_background(f"Página {page:02d}")
                y = height - 112
            if paragraph.startswith("- "):
                c.setFillColor(colors.HexColor("#18a999"))
                c.circle(94, y + 4, 2.5, fill=1, stroke=0)
                y = _draw_wrapped(c, paragraph[2:], 108, y, width - 176, "Helvetica", 10.5, 14, colors.HexColor("#25324a"), max_lines=6) - 4
            elif paragraph.startswith("|"):
                y = _draw_wrapped(c, paragraph.replace("|", "  "), 96, y, width - 166, "Helvetica", 8.5, 12, colors.HexColor("#25324a"), max_lines=5) - 2
            else:
                y = _draw_wrapped(c, paragraph, 96, y, width - 166, "Helvetica", 10.5, 15, colors.HexColor("#25324a"), max_lines=10) - 6
        c.showPage()
        page += 1

    c.save()
    return path


def export_pdf(report: Report) -> Path:
    if report.template_used in ("Motor propio OPERIA", "CPython AST revisado"):
        return _export_native_pdf(report)
    path = new_export_path(f"report-{report.id}.pdf")
    case = isolated_case_snapshot(report.case)
    styles = _report_styles()
    story: list = []
    total_damage = sum(item.amount for item in case.damage_items)
    legal_probatory = legal_probatory_assessment(case)
    exportability = report_exportability_assessment(case, legal_probatory, report=report)

    story.append(_executive_cover(report, case, exportability, total_damage, styles))
    story.append(Spacer(1, 18))
    story.append(
        _key_value_table(
            [
                ("Estado del caso", _professional_status(report, case)),
                ("Jurisdicci\u00f3n", case.jurisdiction),
                ("Responsable", _responsible_text(report)),
                ("Cliente / damnificado", case.client_or_victim or ""),
                ("Autor del informe", _visible_author(report) or "Analista no informado"),
                ("Versi\u00f3n", report.version),
                ("Estado del informe", _professional_status(report, case)),
                ("Nivel de exportabilidad", exportability["level"]),
                ("Puntaje de revisi\u00f3n", str(report.review_score or 0)),
            ],
            styles,
        )
    )
    story.append(Spacer(1, 16))
    story.append(
        _metrics_table(
            [
                ("Participantes", len(case.participants)),
                ("Evidencias", len(case.evidence)),
                ("Custodia", len(case.actions)),
                ("Hallazgos", len(case.findings)),
                ("Mapeos legales", len(case.legal_mappings)),
                ("Da\u00f1o estimado", f"{total_damage:,.2f} ARS"),
            ],
            styles,
        )
    )
    story.append(Spacer(1, 10))
    story.append(_executive_note("Documento preparado para lectura ejecutiva y revisión pericial. Las secciones, anexos y metadatos se presentan según la información efectivamente cargada y validada en el caso.", styles))
    story.append(PageBreak())

    _section(story, "1. Resumen ejecutivo", styles)
    executive_summary = build_smart_executive_summary(case, legal_probatory, exportability)
    _append_markdown_content(story, executive_summary, styles)
    story.append(Spacer(1, 6))
    story.append(Paragraph("Nivel de exportabilidad del informe", styles["subheading"]))
    _append_markdown_content(story, exportability_markdown(exportability), styles)

    _append_challenge_2026_pdf_sections(story, case, legal_probatory, styles)

    investigative_content = _extract_named_sections(
        report.markdown_content,
        [
            "Tecnicas usadas, herramientas forenses y metodologia",
            "T\u00e9cnicas usadas, herramientas forenses y metodolog\u00eda",
            "Identificacion, categorizacion y preservacion de evidencia",
            "Identificaci\u00f3n, categorizaci\u00f3n y preservaci\u00f3n de evidencia",
            "Documentacion investigada",
            "Documentaci\u00f3n investigada",
            "Evidencia de analisis y trazabilidad del hallazgo",
            "Evidencia de an\u00e1lisis y trazabilidad del hallazgo",
        ],
    )
    if investigative_content:
        _section(story, "T\u00e9cnicas, documentaci\u00f3n y trazabilidad probatoria", styles)
        _append_markdown_content(story, investigative_content, styles)

    if case.participants:
        _section(story, "2. Personas y entidades involucradas", styles)
        story.append(
            _data_table(
                ["Rol", "Tipo", "Nombre", "Documento / CUIT", "Relaci\u00f3n", "Sensibilidad"],
                [
                    [
                        _display_value(item.participant_type),
                        _display_value(item.person_type),
                        _display_name(item.display_name),
                        item.document_number or item.tax_id or "-",
                        _display_value(item.relationship_to_case or item.role_description or "-"),
                        _display_value(item.sensitivity_level),
                    ]
                    for item in case.participants
                ],
                styles,
                [70, 62, 110, 84, 126, 72],
            )
        )

    if case.evidence:
        _section(story, "3. Evidencia digital y control de integridad", styles)
        story.append(
            _data_table(
                ["ID", "Tipo", "Descripci\u00f3n", "SHA-256", "MD5", "Integridad"],
                [
                    [
                        str(item.id),
                        item.evidence_type,
                        item.description,
                        _short_hash(item.sha256_hash),
                        _short_hash(item.md5_hash or ""),
                        item.integrity_status,
                    ]
                    for item in case.evidence
                ],
                styles,
                [28, 82, 160, 118, 70, 66],
            )
        )

    if case.actions:
        _section(story, "4. Cadena de custodia", styles)
        story.append(
            _data_table(
                ["Fecha", "Acci\u00f3n", "Evidencia", "Responsable", "Detalle", "Hash posterior"],
                [
                    [
                        _format_date(item.action_at),
                        item.action_type,
                        str(item.evidence_id or "-"),
                        item.user_responsible or "-",
                        item.detailed_description,
                        _short_hash(item.hash_after or item.hash_before or ""),
                    ]
                    for item in case.actions
                ],
                styles,
                [68, 62, 50, 86, 182, 76],
            )
        )

    if case.findings:
        _section(story, "5. Hallazgos y limitaciones", styles)
        story.append(
            _data_table(
                ["Hallazgo", "Confianza", "Soporte", "Limitaciones"],
                [
                    [
                        item.title,
                        item.confidence_level,
                        item.supporting_evidence or item.technical_finding,
                        item.weaknesses_or_limitations or "-",
                    ]
                    for item in case.findings
                ],
                styles,
                [128, 68, 176, 152],
            )
        )
    if case.hypotheses:
        if not case.findings:
            _section(story, "5. Hip\u00f3tesis investigativas", styles)
        story.append(Spacer(1, 8))
        story.append(
            _data_table(
                ["Hip\u00f3tesis", "Soporte", "Conclusi\u00f3n provisional"],
                [[item.title, item.support_degree, item.provisional_conclusion or item.description] for item in case.hypotheses],
                styles,
                [150, 80, 294],
            )
        )

    legal_content = _extract_named_sections(
        report.markdown_content,
        [
            "Encuadre jur\u00eddico argentino",
            "Fundamentacion legal argentina con articulos citados",
            "Fundamentaci\u00f3n legal argentina con art\u00edculos citados",
            "Fundamentacion penal y legal aplicada",
            "Fundamentaci\u00f3n penal y legal aplicada",
        ],
    )
    approved_mappings = [
        item for item in case.legal_mappings if getattr(item, "decision_status", "sugerido") != "descartado"
    ]
    if approved_mappings or legal_content:
        _section(story, "6. Encuadre legal argentino", styles)
    if approved_mappings:
        story.append(
            _data_table(
                ["\u00c1rea", "Norma / art\u00edculo", "An\u00e1lisis", "Decisi\u00f3n", "Confianza"],
                [
                    [
                        item.legal_area,
                        item.norm_reference,
                        item.analysis,
                        getattr(item, "decision_status", "") or item.verification_status,
                        item.confidence_level,
                    ]
                    for item in approved_mappings
                ],
                styles,
                [78, 116, 190, 90, 50],
            )
        )
        linked_rows = []
        for mapping in approved_mappings:
            for link in getattr(mapping, "evidence_links", []):
                if not getattr(link, "evidence", None):
                    continue
                linked_rows.append(
                    [
                        mapping.norm_reference,
                        link.legal_element or "",
                        f"#{link.evidence.id} {link.evidence.description}",
                        getattr(link, "element_status", "pendiente"),
                        link.probatory_strength,
                        getattr(link, "contradiction_risk", "bajo"),
                        getattr(link, "recommended_measure", "") or link.link_rationale or "",
                    ]
                )
        if linked_rows:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Matriz legal-probatoria", styles["subheading"]))
            story.append(
                _data_table(
                    ["Norma", "Elemento", "Evidencia", "Estado", "Fuerza", "Riesgo", "Medida"],
                    linked_rows,
                    styles,
                    [76, 86, 128, 56, 48, 46, 138],
                )
            )
    if legal_content and not approved_mappings:
        _append_markdown_content(story, legal_content, styles)
    sufficiency_content = _extract_named_sections(
        report.markdown_content,
        ["Matriz legal-probatoria aplicada", "Sem\u00e1foro de suficiencia legal"],
    )
    if sufficiency_content:
        story.append(Spacer(1, 8))
        story.append(Paragraph("Matriz legal-probatoria aplicada y sem\u00e1foro de suficiencia legal", styles["subheading"]))
        _append_markdown_content(story, sufficiency_content, styles)

    analysis_content = _extract_named_sections(
        report.markdown_content,
        [
            "Hip\u00f3tesis investigativas",
            "Hip\u00f3tesis y an\u00e1lisis del caso",
            "Hipotesis y analisis del caso",
            "Evidencia que sostiene o limita la teor\u00eda del caso",
            "Evidencia que sostiene o limita la teoria del caso",
        ],
    )
    if analysis_content:
        _section(story, "7. Hip\u00f3tesis y an\u00e1lisis del caso", styles)
        _append_markdown_content(story, analysis_content, styles)

    if getattr(case, "deepfake_analyses", None):
        _section(story, "7.1 Autenticaci\u00f3n de medios y an\u00e1lisis deepfake", styles)
        story.append(
            _data_table(
                ["Medio", "Evidencia / hash", "Herramienta", "Clasificaci\u00f3n", "Hallazgos y l\u00edmites"],
                [
                    [
                        item.media_type,
                        f"#{item.evidence_id or '-'} / {item.evidence_hash or 'pendiente'}",
                        item.tool_used or "pendiente",
                        item.classification,
                        f"{item.findings or 'Sin hallazgos concluyentes cargados.'} L\u00edmites: {item.limitations or 'requiere pericia especializada.'}",
                    ]
                    for item in sorted(case.deepfake_analyses, key=lambda row: row.created_at, reverse=True)
                ],
                styles,
                [58, 112, 92, 88, 174],
            )
        )
        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                _html(
                    "Conclusi\u00f3n prudente: la autenticaci\u00f3n de medios es preliminar y no permite afirmar de forma concluyente que el archivo sea deepfake, aut\u00e9ntico o atribuible a una persona sin pericia especializada, preservaci\u00f3n del original y corroboraci\u00f3n independiente."
                ),
                styles["body"],
            )
        )

    blockchain_transactions = [item for item in (getattr(case, "blockchain_transactions", []) or []) if getattr(item, "include_in_report", True)]
    if not blockchain_transactions:
        blockchain_transactions = list(getattr(case, "blockchain_transactions", []) or [])
    blockchain_transactions = sorted(blockchain_transactions, key=lambda item: (getattr(item, "route_order", 0) or 999999, item.created_at))
    if blockchain_transactions:
        _section(story, "7.2 An\u00e1lisis forense de blockchain - Blockchain Investigation Lab", styles)
        story.append(
            Paragraph(
                _html(
                    "Se documentan movimientos de criptoactivos, wallets, exploradores consultados, evidencia vinculada y "
                    "ruta economica probable. La trazabilidad blockchain orienta el recorrido patrimonial, pero no acredita "
                    "por si sola identidad personal ni titularidad efectiva de una wallet."
                ),
                styles["body"],
            )
        )
        story.append(Spacer(1, 6))
        story.append(
            _data_table(
                ["Red", "Hash / fuente", "Ruta", "Monto", "Soporte", "Hipotesis, suficiencia y limites"],
                [
                    [
                        item.blockchain,
                        f"{item.tx_hash}\n{item.explorer_url or 'URL no informada'}",
                        f"{item.source_wallet or 'pendiente'} -> {item.destination_wallet or 'pendiente'}",
                        f"{item.amount} {item.token} / valor aprox. {item.approximate_value:.2f}",
                        f"Evidencia #{item.evidence_id or '-'}\nHash: {item.evidence_hash or 'pendiente'}\nHerramienta: {item.tool_used or 'pendiente'}\nClasificacion: {item.wallet_classification}",
                        f"{item.route_hypothesis or 'Pendiente.'} Suficiencia: {_blockchain_pdf_readiness(item)[0]} ({_blockchain_pdf_readiness(item)[1]}/100). Limites: {item.limitations or 'No atribuir identidad personal solo por wallet.'}",
                    ]
                    for item in blockchain_transactions
                ],
                styles,
                [52, 104, 112, 72, 76, 108],
            )
        )
        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                _html("Ruta multi-salto observada:"),
                styles["subheading"],
            )
        )
        story.append(
            _data_table(
                ["Salto", "Origen", "Destino", "Monto", "Hash"],
                _blockchain_pdf_multihop_rows(blockchain_transactions),
                styles,
                [42, 128, 128, 82, 144],
            )
        )
        story.append(Spacer(1, 6))
        story.append(Paragraph(_html("Anexo blockchain profesional:"), styles["subheading"]))
        story.append(
            _data_table(
                ["Registro", "Acta / fuente", "Preservación", "Checklist", "Uso pericial"],
                _blockchain_pdf_annex_rows(blockchain_transactions),
                styles,
                [50, 130, 128, 108, 108],
            )
        )
        story.append(Spacer(1, 6))
        story.append(Paragraph(_html("Puntos críticos y requerimientos sugeridos:"), styles["subheading"]))
        story.append(
            _data_table(
                ["Registro", "Punto crítico", "Medida sugerida"],
                _blockchain_pdf_critical_rows(blockchain_transactions),
                styles,
                [62, 210, 252],
            )
        )
        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                _html(
                    "L\u00edmite de atribuci\u00f3n: una wallet o transacci\u00f3n permite orientar trazabilidad econ\u00f3mica, pero no identifica por s\u00ed sola a una persona sin evidencia externa corroborante, informaci\u00f3n de plataforma, comunicaciones, cuentas, IPs u otros registros."
                ),
                styles["body"],
            )
        )

    if getattr(case, "iot_devices", None):
        _section(story, "7.3 An\u00e1lisis forense de dispositivos IoT", styles)
        story.append(
            _data_table(
                ["Dispositivo", "Red", "Evidencia", "Checklist / timeline", "Hallazgos y l\u00edmites"],
                [
                    [
                        f"{item.device_type} / {item.manufacturer or ''} {item.model or ''}",
                        f"IP int. {item.internal_ip or '-'} / IP ext. {item.external_ip or '-'} / MAC {item.mac or '-'}",
                        f"#{item.evidence_id or '-'} / {item.evidence_hash or 'pendiente'}",
                        f"Checklist: {_iot_pdf_checklist_summary(item.checklist or '')}\nTimeline: identificaci\u00f3n, preservaci\u00f3n, observaci\u00f3n t\u00e9cnica, hallazgo, norma e informe.",
                        f"Logs: {item.logs_available or 'pendiente'}; puertos: {item.open_ports or 'pendiente'}; vulnerabilidades: {item.known_vulnerabilities or 'pendiente'}. L\u00edmites: {item.limitations or 'No atribuir por IP, MAC o puerto sin corroboraci\u00f3n.'}",
                    ]
                    for item in sorted(case.iot_devices, key=lambda row: row.created_at, reverse=True)
                ],
                styles,
                [86, 116, 90, 108, 124],
            )
        )
        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                _html(
                    "L\u00edmite de atribuci\u00f3n: un dispositivo IoT, IP, MAC, puerto o firmware observado no identifica por s\u00ed solo a una persona ni acredita acceso ileg\u00edtimo sin logs, autorizaciones, timeline y evidencia externa corroborante."
                ),
                styles["body"],
            )
        )

    conclusion_content = _extract_named_sections(
        report.markdown_content,
        [
            "Conclusi\u00f3n jur\u00eddicamente defendible",
            "Conclusiones profundas y prudentes",
            "Conclusiones t\u00e9cnicas",
            "Conclusi\u00f3n t\u00e9cnico-jur\u00eddica",
            "Conclusion tecnico-juridica",
        ],
    )
    _section(story, "8. Conclusi\u00f3n t\u00e9cnico-jur\u00eddica", styles)
    if conclusion_content:
        _append_markdown_content(story, conclusion_content, styles)
    else:
        story.append(Paragraph(_html(build_matrix_based_conclusion(case)), styles["body"]))
        story.append(Spacer(1, 6))
        story.append(Paragraph(_html(build_attribution_conclusion(case)), styles["body"]))

    recommendations_content = _extract_named_sections(
        report.markdown_content,
        ["Recomendaciones procesales autom\u00e1ticas", "Recomendaciones"],
    )
    if recommendations_content:
        _section(story, "9. Recomendaciones procesales", styles)
        _append_markdown_content(story, recommendations_content, styles)

    if getattr(case, "rubric_evaluations", None):
        from operia.rubrics.service import evaluation_payload

        latest_rubric = sorted(case.rubric_evaluations, key=lambda item: item.created_at, reverse=True)[0]
        rubric = evaluation_payload(latest_rubric)
        _section(story, "10. Cumplimiento de r\u00fabrica del informe", styles)
        story.append(
            _key_value_table(
                [
                    ("R\u00fabrica", rubric["rubric_title"]),
                    ("Score", f"{rubric['total_score']} / {rubric['max_score']} ({rubric['percentage']}%)"),
                    ("Estado", rubric["status"]),
                    ("Observaci\u00f3n", rubric["summary"]),
                ],
                styles,
            )
        )
        story.append(Spacer(1, 8))
        story.append(
            _data_table(
                ["Secci\u00f3n", "Estado", "Puntaje", "Faltantes / observaciones"],
                [
                    [
                        f"{section['key']}. {section['label']}",
                        section["status"],
                        f"{section['estimated_score']} / {section['max_score']}",
                        "; ".join(section.get("missing") or section.get("recommendations") or ["Sin observaciones cr\u00edticas"]),
                    ]
                    for section in rubric["sections"]
                ],
                styles,
                [150, 62, 58, 254],
            )
        )

    _section(story, "11. Anexo de exportaci\u00f3n", styles)
    story.append(
        Paragraph(
            _html(
                _export_metadata_text(report)
            ),
            styles["body"],
        )
    )
    annexes = report_annexes_payload(report, legal_probatory)
    if annexes["items"]:
        story.append(Spacer(1, 8))
        story.append(Paragraph("Anexos reales incorporados", styles["subheading"]))
        story.append(
            _data_table(
                ["Tipo", "T\u00edtulo", "Hash / referencia", "Detalle"],
                [
                    [
                        item.get("type", ""),
                        item.get("title", ""),
                        item.get("sha256") or item.get("hash_after") or item.get("version") or "",
                        item.get("integrity_status") or item.get("detail") or f"Filas: {item.get('rows', '-')}",
                    ]
                    for item in annexes["items"]
                ],
                styles,
                [82, 190, 120, 132],
            )
        )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=54,
        bottomMargin=48,
        title=normalize_spanish(report.title),
        author=normalize_spanish(_visible_author(report) or "Analista no informado"),
    )
    doc.build(story, onFirstPage=_report_page_footer, onLaterPages=_report_page_footer)
    return path


def _report_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "cover_title": ParagraphStyle(
            "OperiaCoverTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=27,
            leading=32,
            textColor=colors.white,
            spaceAfter=8,
        ),
        "cover_subtitle": ParagraphStyle(
            "OperiaCoverSubtitle",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=11,
            leading=16,
            textColor=colors.HexColor("#dbeafe"),
        ),
        "cover_meta": ParagraphStyle(
            "OperiaCoverMeta",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#a7f3d0"),
        ),
        "label": ParagraphStyle(
            "OperiaLabel",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#0f766e"),
            spaceAfter=8,
        ),
        "subtitle": ParagraphStyle(
            "OperiaSubtitle",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#475569"),
        ),
        "heading": ParagraphStyle(
            "OperiaHeading",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            spaceBefore=18,
            spaceAfter=10,
        ),
        "subheading": ParagraphStyle(
            "OperiaSubheading",
            parent=base["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#155e75"),
            spaceBefore=8,
            spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "OperiaBody",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=13.8,
            textColor=colors.HexColor("#25324a"),
            spaceAfter=6,
            alignment=TA_JUSTIFY,
        ),
        "muted": ParagraphStyle(
            "OperiaMuted",
            parent=base["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor("#64748b"),
        ),
        "table": ParagraphStyle(
            "OperiaTable",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.4,
            leading=9.2,
            textColor=colors.HexColor("#243041"),
        ),
        "table_header": ParagraphStyle(
            "OperiaTableHeader",
            parent=base["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=7.2,
            leading=8.8,
            textColor=colors.white,
        ),
    }


def _public_text(value: object) -> str:
    text = normalize_spanish("" if value is None else str(value))
    text = text.replace("enriquecido_con_ia", "En Proceso Investigativo")
    text = text.replace("manifiesto_exportacion", "manifiesto de exportación")
    text = text.replace("manifesto_exportacion", "manifiesto de exportación")
    text = re.sub(r"\bBORRADOR_CONTROLADO\b", "Borrador controlado", text, flags=re.IGNORECASE)
    text = re.sub(r"\bVersion\b", "Versión", text)
    text = re.sub(r"\bgeneracion\b", "generación", text)
    text = re.sub(r"\btecnicos\b", "técnicos", text)
    text = re.sub(r"\bexportacion\b", "exportación", text)
    text = re.sub(r"\bdisposicion\b", "disposición", text)
    text = re.sub(r"\bDisposicion\b", "Disposición", text)
    text = re.sub(r"\baccion\b", "acción", text)
    text = re.sub(r"\bacciones\b", "acciones", text)
    text = re.sub(r"\bpreservacion\b", "preservación", text)
    text = re.sub(r"\bidentificacion\b", "identificación", text)
    text = re.sub(r"\binformacion\b", "información", text)
    return text


def _html(value: object) -> str:
    return escape(_public_text("" if value is None else str(value))).replace("\n", "<br/>")


def _paragraph(value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(_html(value), style)


def _executive_cover(report: Report, case, exportability: dict, total_damage: float, styles: dict[str, ParagraphStyle]):
    status_tone = "INFORME FINAL" if exportability.get("can_export_final") else "INFORME PRELIMINAR"
    subtitle = f"Caso {case.case_code} | {case.incident_type} | Jurisdicción {case.jurisdiction}"
    metadata = [
        f"Estado: {_professional_status(report, case)}",
        f"Exportabilidad: {exportability.get('level', '-')}",
        f"Revisión: {report.review_score or 0}/100",
        f"Daño estimado: {total_damage:,.2f} ARS",
    ]
    data = [
        [Paragraph("INFORME FORENSE PROFESIONAL", styles["cover_meta"]), Paragraph(status_tone, styles["cover_meta"])],
        [Paragraph(_html(report.title), styles["cover_title"]), ""],
        [Paragraph(_html(subtitle), styles["cover_subtitle"]), ""],
        [Paragraph(_html("  |  ".join(metadata)), styles["cover_subtitle"]), ""],
    ]
    table = Table(data, colWidths=[390, 134], rowHeights=[26, 98, 34, 42])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0b172a")),
                ("SPAN", (0, 1), (-1, 1)),
                ("SPAN", (0, 2), (-1, 2)),
                ("SPAN", (0, 3), (-1, 3)),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#0f766e")),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#1e3a8a")),
                ("LEFTPADDING", (0, 0), (-1, -1), 22),
                ("RIGHTPADDING", (0, 0), (-1, -1), 22),
                ("TOPPADDING", (0, 0), (-1, -1), 12),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    return table


def _executive_note(text: str, styles: dict[str, ParagraphStyle]):
    table = Table([[Paragraph(_html(text), styles["muted"])]], colWidths=[524])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f0fdfa")),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#99f6e4")),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
            ]
        )
    )
    return table


def _section(story: list, title: str, styles: dict[str, ParagraphStyle]) -> None:
    table = Table([[Paragraph(_html(title), styles["heading"])]], colWidths=[524])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
                ("LINEBEFORE", (0, 0), (0, -1), 4, colors.HexColor("#0f766e")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(table)


def _blockchain_pdf_readiness(tx) -> tuple[str, int, list[str]]:
    checks = [
        ("hash de transaccion", bool(getattr(tx, "tx_hash", ""))),
        ("URL o fuente preservada", bool(getattr(tx, "explorer_url", ""))),
        ("wallet origen y destino", bool(getattr(tx, "source_wallet", "") and getattr(tx, "destination_wallet", ""))),
        ("monto y token", bool((getattr(tx, "amount", 0) or 0) > 0 and getattr(tx, "token", ""))),
        ("fecha/hora observada", bool(getattr(tx, "tx_datetime", ""))),
        ("evidencia vinculada", bool(getattr(tx, "evidence_id", None))),
        ("hipotesis de ruta", bool(getattr(tx, "route_hypothesis", ""))),
        ("hallazgo creado", bool(getattr(tx, "finding_id", None))),
        ("normas aplicadas", bool(getattr(tx, "legal_mapping_ids", ""))),
    ]
    ok_count = sum(1 for _, ok in checks if ok)
    score = round((ok_count / len(checks)) * 100)
    missing = [label for label, ok in checks if not ok]
    if score >= 78:
        return "listo para informe", score, missing
    if score >= 45:
        return "requiere refuerzo", score, missing
    return "insuficiente", score, missing


def _blockchain_pdf_multihop_rows(transactions) -> list[list[str]]:
    ordered = sorted(transactions, key=lambda item: (getattr(item, "route_order", 0) or 999999, item.created_at))
    rows = []
    for index, item in enumerate(ordered, start=1):
        rows.append(
            [
                str(index),
                item.source_wallet or "Pendiente",
                item.destination_wallet or "Pendiente",
                f"{item.amount} {item.token}",
                item.tx_hash or "Pendiente",
            ]
        )
    return rows or [["-", "Sin registros", "-", "-", "-"]]


def _blockchain_pdf_annex_rows(transactions) -> list[list[str]]:
    rows = []
    for item in sorted(transactions, key=lambda tx: (getattr(tx, "route_order", 0) or 999999, tx.created_at)):
        readiness_label, readiness_score, readiness_missing = _blockchain_pdf_readiness(item)
        source = item.explorer_url or "Fuente no informada"
        evidence = f"Evidencia #{item.evidence_id}" if getattr(item, "evidence_id", None) else "Sin evidencia vinculada"
        evidence_hash = item.evidence_hash or "Hash pendiente"
        act = f"{item.tx_datetime or 'Fecha/hora pendiente'}\n{item.tool_used or 'Herramienta pendiente'}\n{source}"
        preservation = (
            f"{evidence}\nHash: {evidence_hash}\n"
            f"Origen: {item.source_wallet or 'Pendiente'}\nDestino: {item.destination_wallet or 'Pendiente'}"
        )
        checklist = (
            f"{readiness_label} ({readiness_score}/100)\n"
            f"Faltantes: {', '.join(readiness_missing[:4]) if readiness_missing else 'sin faltantes criticos'}"
        )
        use = (
            "Trazabilidad economica y medidas investigativas. "
            "No acredita identidad personal sin corroboracion externa."
        )
        rows.append([f"#{item.id}", act, preservation, checklist, use])
    return rows or [["-", "Sin registros", "-", "-", "-"]]


def _blockchain_pdf_critical_rows(transactions) -> list[list[str]]:
    rows = []
    for item in sorted(transactions, key=lambda tx: (getattr(tx, "route_order", 0) or 999999, tx.created_at)):
        missing = []
        if not getattr(item, "evidence_id", None):
            missing.append("evidencia")
        if not getattr(item, "explorer_url", ""):
            missing.append("fuente/URL")
        if not getattr(item, "tx_datetime", ""):
            missing.append("fecha/hora")
        if not getattr(item, "finding_id", None):
            missing.append("hallazgo")
        if not getattr(item, "legal_mapping_ids", ""):
            missing.append("normas")
        classification = (getattr(item, "wallet_classification", "") or "").lower()
        critical = ", ".join(missing) if missing else "Sin faltantes críticos de documentación inicial"
        if classification in {"exchange", "mixer", "bridge", "contrato inteligente"}:
            critical = f"{critical}. Punto de servicio/riesgo: {classification}"
        measure = "Vincular evidencia, preservar captura/URL/fecha y correlacionar con comunicaciones, plataforma, IPs, KYC o registros externos antes de atribuir."
        if classification == "exchange":
            measure = "Requerir KYC, titularidad, logs, IPs, dispositivo y movimientos a la plataforma por vía legal."
        elif classification == "mixer":
            measure = "Documentar límite de trazabilidad por mixer y evitar continuidad de fondos sin análisis complementario."
        elif classification == "bridge":
            measure = "Preservar transacciones en red origen y destino, hash, wallet, monto y fecha/hora."
        elif classification == "contrato inteligente":
            measure = "Revisar eventos del contrato, función invocada, dirección y wallets firmantes."
        rows.append([f"#{item.id}", critical, measure])
    return rows or [["-", "Sin registros", "-"]]


def _iot_pdf_checklist_summary(raw_checklist: str) -> str:
    labels = {
        "physical_identification": "identificación física",
        "photo_capture": "fotografía/captura",
        "hash_evidence": "hash de evidencia",
        "firmware": "firmware",
        "configuration": "configuración",
        "logs": "logs",
        "ports": "puertos",
        "default_credentials": "credenciales por defecto",
        "network_connection": "conexión a red",
        "timeline": "línea temporal",
        "state_preservation": "preservación de estado",
        "alteration_risk": "riesgo de alteración",
    }
    try:
        parsed = json.loads(raw_checklist or "{}")
    except json.JSONDecodeError:
        return "checklist no estructurado"
    done = [labels.get(key, key) for key, value in parsed.items() if isinstance(value, dict) and value.get("ok")]
    return ", ".join(done) if done else "sin controles marcados"


def _warning_box(text: str):
    style = _report_styles()["table"]
    table = Table([[Paragraph(_html(text), style)]], colWidths=[524])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fff7ed")),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#fed7aa")),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    return table


def _key_value_table(rows: list[tuple[str, object]], styles: dict[str, ParagraphStyle]):
    visible_rows = [(label, value) for label, value in rows if _public_text(value).strip()]
    data = [[_paragraph(label, styles["table_header"]), _paragraph(value, styles["table"])] for label, value in visible_rows]
    table = Table(data, colWidths=[150, 374])
    table.setStyle(_base_table_style(header_column=True))
    return table


def _metrics_table(metrics: list[tuple[str, object]], styles: dict[str, ParagraphStyle]):
    cells = []
    for label, value in metrics:
        cells.append(Paragraph(f"<font size='12'><b>{_html(value)}</b></font><br/><font color='#64748b'>{_html(label)}</font>", styles["table"]))
    data = [cells[:3], cells[3:]]
    table = Table(data, colWidths=[174, 175, 175])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#ffffff")),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#cbd5e1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
                ("LEFTPADDING", (0, 0), (-1, -1), 14),
                ("RIGHTPADDING", (0, 0), (-1, -1), 14),
                ("TOPPADDING", (0, 0), (-1, -1), 12),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
            ]
        )
    )
    return table


def _data_table(headers: list[str], rows: list[list[object]], styles: dict[str, ParagraphStyle], col_widths: list[int]):
    data = [[_paragraph(header, styles["table_header"]) for header in headers]]
    data.extend([[_paragraph(cell, styles["table"]) for cell in row] for row in rows])
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(_base_table_style())
    return table


def _append_challenge_2026_pdf_sections(story: list, case, legal_probatory: dict, styles: dict[str, ParagraphStyle]) -> None:
    _section(story, "Written Legal Evidence Report 2026", styles)
    story.append(Paragraph("A. Technical Approach & Investigative Techniques Used", styles["subheading"]))
    story.append(
        _data_table(
            ["Herramienta / metodologia", "Motivo", "Tareas realizadas", "Limites"],
            _challenge_pdf_technical_rows(case),
            styles,
            [124, 124, 150, 126],
        )
    )
    story.append(Spacer(1, 8))
    story.append(Paragraph("B. Evidence Identification, Categorization & Handling", styles["subheading"]))
    story.append(
        _data_table(
            ["Evidencia", "Tipo", "Preservacion", "Hash", "Custodia"],
            _challenge_pdf_evidence_rows(case),
            styles,
            [116, 70, 116, 108, 114],
        )
    )
    story.append(Spacer(1, 8))
    story.append(Paragraph("C. Legal Authority & Investigative Requests", styles["subheading"]))
    story.append(
        _data_table(
            ["Autoridad legal", "Solicitud", "Fundamento", "Estado", "Documentacion"],
            _challenge_pdf_legal_rows(case, legal_probatory),
            styles,
            [100, 118, 132, 66, 108],
        )
    )
    story.append(Spacer(1, 8))
    story.append(Paragraph("D. Summary Analysis of Evidence & Investigative Findings", styles["subheading"]))
    story.append(
        _data_table(
            ["Evidencia / hallazgo", "Relevancia investigativa", "Significancia legal", "Atribucion / limites"],
            _challenge_pdf_finding_rows(case),
            styles,
            [130, 134, 128, 132],
        )
    )
    story.append(Spacer(1, 8))
    story.append(Paragraph("Final Investigative Summary", styles["subheading"]))
    for line in _challenge_pdf_summary_lines(case):
        story.append(Paragraph(_html(line), styles["body"]))


def _challenge_pdf_technical_rows(case) -> list[list[object]]:
    rows: list[list[object]] = []
    tools = []
    for evidence in case.evidence:
        if evidence.tool_used:
            tools.append(evidence.tool_used)
    for action in case.actions:
        if action.tool_used:
            tools.append(action.tool_used)
    for finding in case.findings:
        if finding.tools_used:
            tools.append(finding.tools_used)
    for tool in dict.fromkeys(_public_text(item).strip() for item in tools if _public_text(item).strip()):
        rows.append(
            [
                tool,
                "Relacion directa con evidencia registrada.",
                "Preservacion, verificacion, analisis o documentacion.",
                "No sustituye autorizacion legal ni revision pericial.",
            ]
        )
    if getattr(case, "deepfake_analyses", None):
        rows.append(
            [
                "Autenticacion de medios",
                "Evaluar medio sospechoso.",
                "Origen, hash, metadatos, inconsistencias y comparacion.",
                "Resultado preliminar; no afirma autenticidad ni autoria.",
            ]
        )
    if getattr(case, "blockchain_transactions", None):
        rows.append(
            [
                "Analisis forense de blockchain",
                "Documentar ruta economica.",
                "Wallets, hash de transaccion, red, monto, URL aportada e hipotesis.",
                "Una wallet no identifica por si sola a una persona.",
            ]
        )
    if getattr(case, "iot_devices", None):
        rows.append(
            [
                "Analisis de dispositivos conectados",
                "Documentar dispositivo y estado observable.",
                "Identificacion, firmware, red, logs, checklist y limitaciones.",
                "IP, MAC o puerto requieren corroboracion externa.",
            ]
        )
    if not rows:
        rows.append(
            [
                "Metodologia documental forense",
                "Estado inicial del caso.",
                "Ordenamiento de evidencia, custodia, hallazgos y matriz legal.",
                "Requiere completar evidencia y revision profesional.",
            ]
        )
    return rows


def _challenge_pdf_evidence_rows(case) -> list[list[object]]:
    rows = []
    for evidence in case.evidence:
        custody = [
            action.action_type
            for action in case.actions
            if getattr(action, "evidence_id", None) == evidence.id
        ]
        rows.append(
            [
                f"#{evidence.id} {evidence.description}",
                evidence.evidence_type or "digital",
                evidence.preservation_method or evidence.integrity_status or "pendiente",
                _short_hash(evidence.sha256_hash or ""),
                ", ".join(custody) if custody else "pendiente de ampliar",
            ]
        )
    if not rows:
        rows.append(["Sin evidencia individual cargada", "Pendiente", "Pendiente", "-", "Pendiente"])
    return rows


def _challenge_pdf_legal_rows(case, legal_probatory: dict) -> list[list[object]]:
    rows = []
    approved = [
        item for item in case.legal_mappings if getattr(item, "decision_status", "sugerido") != "descartado"
    ]
    for mapping in approved:
        rows.append(
            [
                mapping.norm_reference,
                "Corroboracion y autorizaciones segun hechos.",
                mapping.applicability_to_case or mapping.analysis or "pendiente de ampliar",
                getattr(mapping, "decision_status", "") or mapping.verification_status,
                mapping.probatory_limits or "fuente registrada en biblioteca legal",
            ]
        )
    if not rows:
        measures = legal_probatory.get("recommended_measures") if isinstance(legal_probatory, dict) else None
        rows.append(
            [
                "Pendiente de vincular",
                "; ".join(measures[:2]) if measures else "Completar mapeo legal.",
                "No formular encuadre definitivo sin fuente y evidencia.",
                "Pendiente",
                "Registrar autorizacion, oficio o decision profesional.",
            ]
        )
    return rows


def _challenge_pdf_finding_rows(case) -> list[list[object]]:
    rows = []
    for finding in case.findings:
        rows.append(
            [
                finding.title,
                finding.technical_finding or finding.supporting_evidence or "pendiente",
                finding.preliminary_legal_relevance or "pendiente",
                finding.weaknesses_or_limitations or "No atribuir autoria personal sin evidencia suficiente.",
            ]
        )
    if not rows:
        rows.append(
            [
                "Sin hallazgos conclusivos cargados",
                "Requiere analisis adicional.",
                "Sin significancia legal definitiva.",
                "No atribuir identidad ni autoria sin evidencia suficiente.",
            ]
        )
    return rows


def _challenge_pdf_summary_lines(case) -> list[str]:
    major_findings = "; ".join(f.title for f in case.findings[:3]) or "hallazgos mayores pendientes de carga"
    next_steps = []
    if not case.evidence:
        next_steps.append("incorporar evidencia identificable")
    if not any(e.sha256_hash for e in case.evidence):
        next_steps.append("registrar hash")
    if not case.actions:
        next_steps.append("registrar cadena de custodia")
    if not case.legal_mappings:
        next_steps.append("vincular matriz legal-probatoria")
    if not next_steps:
        next_steps.append("mantener revision profesional, anexos y corroboracion externa")
    return [
        f"Importancia general: el informe organiza el incidente {case.incident_type} con evidencia, custodia, hallazgos y encuadre legal disponible.",
        f"Hallazgos mayores: {major_findings}.",
        "Observaciones y limitaciones: las conclusiones son prudentes y no atribuyen autoria personal sin evidencia tecnica y legal suficiente.",
        f"Proximos pasos: {', '.join(next_steps)}.",
    ]


def _base_table_style(header_column: bool = False) -> TableStyle:
    commands = [
        ("BOX", (0, 0), (-1, -1), 0.65, colors.HexColor("#cbd5e1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    if header_column:
        commands.append(("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#12324a")))
        commands.append(("BACKGROUND", (1, 0), (1, -1), colors.HexColor("#ffffff")))
    else:
        commands.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12324a")))
        commands.append(("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]))
    return TableStyle(commands)


def _append_markdown_content(story: list, content: str, styles: dict[str, ParagraphStyle]) -> None:
    for raw_line in normalize_spanish(content).splitlines():
        line = raw_line.strip()
        if not line or line.startswith("| ---"):
            continue
        if _is_internal_line(line) or not _is_public_line(line):
            continue
        if line.startswith("### "):
            title = line[4:].strip()
            if not _is_internal_line(title) and _is_public_line(title):
                story.append(Paragraph(_html(title), styles["subheading"]))
        elif line.startswith("## "):
            title = line[3:].strip()
            if not _is_internal_line(title) and _is_public_line(title):
                story.append(Paragraph(_html(title), styles["subheading"]))
        elif line.startswith("- "):
            story.append(Paragraph(f"\u2022 {_html(_clean_markdown_text(line[2:]))}", styles["body"]))
        elif line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if len(cells) > 1:
                story.append(Paragraph(_html(" | ".join(cells)), styles["table"]))
        else:
            story.append(Paragraph(_html(_clean_markdown_text(line)), styles["body"]))


def _extract_named_sections(markdown: str, titles: list[str]) -> str:
    wanted = {normalize_spanish(title).lower() for title in titles}
    parts = []
    seen = set()
    for title, content in _markdown_sections(markdown):
        normalized_title = normalize_spanish(title).lower()
        normalized_content = normalize_spanish(content).strip()
        fingerprint = re.sub(r"\s+", " ", normalized_content.lower())
        if normalized_title in wanted and fingerprint not in seen:
            seen.add(fingerprint)
            parts.append(content.strip())
    return "\n\n".join(part for part in parts if _meaningful_content(part))


def _meaningful_content(content: str) -> bool:
    text = normalize_spanish(content).strip().lower()
    if not text:
        return False
    empty_markers = [
        "sin encuadre legal",
        "no hay encuadre legal",
        "sin evidencia cargada",
        "sin acciones registradas",
        "sin hallazgos",
        "evaluacion pendiente",
        "evaluación pendiente",
    ]
    return not any(marker in text for marker in empty_markers) and _is_public_line(text)


def _is_internal_line(text: str) -> bool:
    haystack = normalize_spanish(text).lower()
    forbidden = ["ollama", "mistral", "fallback", " ia ", "asistido por ia", "herramientas automatizadas", "enriquecido_con_ia"]
    padded = f" {haystack} "
    return any(term in padded for term in forbidden)


def _is_public_line(text: str) -> bool:
    haystack = normalize_spanish(text).lower()
    return not any(marker in haystack for marker in PUBLIC_FORBIDDEN_MARKERS)


def _display_value(value: object) -> str:
    text = normalize_spanish(str(value or "")).replace("_", " ").strip()
    if not text:
        return ""
    return text[:1].upper() + text[1:]


def _display_name(value: object) -> str:
    text = normalize_spanish(str(value or "")).strip()
    return " ".join(part[:1].upper() + part[1:] for part in text.split())


def _visible_author(report: Report) -> str:
    author = normalize_spanish(report.author or "").strip()
    return "" if author.upper() == "OPERIA" else author


def _responsible_text(report: Report) -> str:
    case = isolated_case_snapshot(report.case)
    assessment = attribution_assessment(case)
    if assessment["sufficient"] and assessment["subject"]:
        return _display_name(assessment["subject"])
    if normalize_spanish(case.responsible or "").strip():
        return _display_name(case.responsible)
    return NO_ATTRIBUTION_TEXT


def _professional_status(report: Report, case) -> str:
    if normalize_spanish(case.status).lower() == "cerrado":
        return "Caso Cerrado"
    if case.findings and case.evidence:
        return "En Proceso Investigativo"
    if case.evidence or case.actions:
        return "En Proceso de Análisis"
    return "Caso en Proceso de Inicio de Investigación"


def _export_metadata_text(report: Report) -> str:
    case = isolated_case_snapshot(report.case)
    rows = [
        f"Identificador del caso: {case.case_code}.",
        f"Versión del informe: {report.version}.",
        f"Fecha y hora de generación: {utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}.",
        f"Autor/analista: {_visible_author(report) or 'No informado'}.",
        f"Estado del caso: {_professional_status(report, case)}.",
        f"Estado documental: {_professional_status(report, case)}.",
        f"Cantidad de evidencias: {len(case.evidence)}.",
        f"Cantidad de acciones de custodia: {len(case.actions)}.",
        "Hash SHA-256, MD5 e identificador de manifiesto: se incorporan al registro técnico de exportación cuando el documento queda emitido.",
    ]
    if report.observations:
        rows.append(f"Observaciones de cierre: {report.observations}.")
    return normalize_spanish("\n".join(rows))


def _professional_conclusion(case) -> str:
    incident = normalize_spanish(case.incident_type or "").lower()
    if "phishing" in incident or "fraude" in incident:
        return (
            "El caso presenta una hip\u00f3tesis preliminar compatible con phishing o fraude digital mediante captaci\u00f3n de v\u00edctimas por "
            "enga\u00f1o, suplantaci\u00f3n o anuncio fraudulento. Desde la normativa argentina, el an\u00e1lisis debe concentrarse en la posible "
            "relaci\u00f3n con el art. 172 del C\u00f3digo Penal, en tanto exige ardid o enga\u00f1o, error de la v\u00edctima, disposici\u00f3n patrimonial, "
            "perjuicio y nexo causal; tambi\u00e9n puede evaluarse el art. 173 inc. 16 si existi\u00f3 manipulaci\u00f3n inform\u00e1tica, el art. 153 bis "
            "si hubo acceso ileg\u00edtimo a cuentas o sistemas, la Ley 25.326 si se comprometieron datos personales y la Ley 24.240 si intervino "
            "una plataforma, proveedor o relaci\u00f3n de consumo. La conclusi\u00f3n no debe cerrar autor\u00eda personal hasta contar con evidencia "
            "t\u00e9cnica preservada, trazabilidad de comunicaciones, registros de pago, identificadores de cuenta, logs, oficios a proveedores y "
            "validaci\u00f3n de cadena de custodia."
        )
    return (
        "La conclusi\u00f3n debe formularse como hip\u00f3tesis t\u00e9cnico-jur\u00eddica preliminar. Para elevar su fuerza probatoria corresponde "
        "conectar cada hecho con evidencia preservada, cada evidencia con un hallazgo verificable y cada hallazgo con una norma aplicable "
        "citada desde la biblioteca legal registrada."
    )


def _short_hash(value: str) -> str:
    value = normalize_spanish(value or "")
    if not value:
        return "-"
    return value if len(value) <= 20 else f"{value[:10]}...{value[-8:]}"


def _format_date(value) -> str:
    if not value:
        return "-"
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M")
    return normalize_spanish(str(value))


def _report_page_footer(c: canvas.Canvas, doc) -> None:
    width, height = A4
    c.saveState()
    c.setFillColor(colors.HexColor("#fbfdff"))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    if doc.page > 1:
        c.setFillColor(colors.HexColor("#0b172a"))
        c.rect(0, height - 34, width, 34, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 8)
        c.drawString(36, height - 21, "OPERIA FORENSIC TRACING")
        c.setFont("Helvetica", 7)
        c.drawRightString(width - 36, height - 21, "Informe forense profesional")
    c.setStrokeColor(colors.HexColor("#d8e4ee"))
    c.line(36, 36, width - 36, 36)
    c.setFont("Helvetica", 7)
    c.setFillColor(colors.HexColor("#64748b"))
    c.drawString(36, 24, "OPERIA Forensic Tracing | Informe forense profesional")
    c.drawRightString(width - 36, 24, f"P\u00e1gina {doc.page}")
    c.restoreState()


def _logo_path() -> Path | None:
    project_root = Path(__file__).resolve().parents[3]
    for name in ("operia-logo-cropped.png", "operia-logo.png"):
        candidate = project_root / "frontend" / "src" / "assets" / name
        if candidate.exists():
            return candidate
    return None


def _fit_text(text: str, width: float, font: str, size: int) -> list[str]:
    lines: list[str] = []
    for raw in text.splitlines() or [""]:
        words = raw.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            test = f"{current} {word}"
            if stringWidth(test, font, size) <= width:
                current = test
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def _draw_wrapped(
    c: canvas.Canvas,
    text: str,
    x: float,
    y: float,
    width: float,
    font: str = "Helvetica",
    size: int = 11,
    leading: int = 15,
    color=colors.HexColor("#25324a"),
    max_lines: int | None = None,
) -> float:
    c.setFont(font, size)
    c.setFillColor(color)
    lines = _fit_text(text, width, font, size)
    if max_lines and len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = f"{lines[-1][: max(0, len(lines[-1]) - 3)]}..."
    for line in lines:
        c.drawString(x, y, line)
        y -= leading
    return y


def _draw_chip(c: canvas.Canvas, x: float, y: float, text: str, fill: str, stroke: str, ink: str) -> None:
    width = stringWidth(text, "Helvetica-Bold", 8) + 18
    c.setFillColor(colors.HexColor(fill))
    c.setStrokeColor(colors.HexColor(stroke))
    c.roundRect(x, y - 13, width, 19, 6, fill=1, stroke=1)
    c.setFillColor(colors.HexColor(ink))
    c.setFont("Helvetica-Bold", 8)
    c.drawString(x + 9, y - 8, text)


def _draw_page_background(c: canvas.Canvas, width: float, height: float) -> None:
    c.setFillColor(colors.white)
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#f5f8fc"))
    c.roundRect(28, 28, width - 56, height - 56, 18, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.roundRect(48, 48, width - 96, height - 96, 18, fill=1, stroke=0)
    c.setStrokeColor(colors.HexColor("#dbe5f1"))
    c.setLineWidth(0.8)
    c.roundRect(48, 48, width - 96, height - 96, 18, fill=0, stroke=1)
    c.setStrokeColor(colors.HexColor("#18a999"))
    c.setLineWidth(2)
    c.line(70, height - 72, width - 70, height - 72)


def _draw_header(c: canvas.Canvas, presentation: Presentation, width: float, height: float, page_number: int) -> None:
    logo = _logo_path()
    if logo:
        c.drawImage(ImageReader(str(logo)), 66, height - 64, width=118, height=45, preserveAspectRatio=True, mask="auto")
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica-Bold", 9)
    c.drawRightString(width - 66, height - 45, normalize_spanish(presentation.presentation_type).upper())
    c.setFont("Helvetica", 8)
    c.setFillColor(colors.HexColor("#667085"))
    c.drawRightString(width - 66, height - 58, f"Página {page_number:02d}")


def _draw_metric(c: canvas.Canvas, x: float, y: float, label: str, value: str, color: str) -> None:
    c.setFillColor(colors.HexColor("#ffffff"))
    c.setStrokeColor(colors.HexColor("#dbe5f1"))
    c.roundRect(x, y, 145, 52, 10, fill=1, stroke=1)
    c.setFillColor(colors.HexColor(color))
    c.circle(x + 17, y + 27, 5, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#64748b"))
    c.setFont("Helvetica-Bold", 7)
    c.drawString(x + 31, y + 32, label.upper())
    c.setFillColor(colors.HexColor("#101828"))
    c.setFont("Helvetica-Bold", 13)
    c.drawString(x + 31, y + 14, value)


def _premium_bg(c: canvas.Canvas, width: float, height: float) -> None:
    c.setFillColor(colors.HexColor("#fbfdff"))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#f2f7fb"))
    c.roundRect(24, 24, width - 48, height - 48, 22, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.roundRect(42, 42, width - 84, height - 84, 18, fill=1, stroke=0)
    c.setStrokeColor(colors.HexColor("#d8e4ee"))
    c.setLineWidth(0.8)
    c.roundRect(42, 42, width - 84, height - 84, 18, fill=0, stroke=1)
    c.setFillColor(colors.HexColor("#e8f7f3"))
    c.rect(42, height - 74, width - 84, 32, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#0f766e"))
    c.rect(42, height - 76, width - 84, 3, fill=1, stroke=0)


def _premium_header(c: canvas.Canvas, presentation: Presentation, width: float, height: float, page_number: int) -> None:
    logo = _logo_path()
    if logo:
        c.drawImage(ImageReader(str(logo)), 62, height - 70, width=128, height=44, preserveAspectRatio=True, mask="auto")
    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(colors.HexColor("#0f766e"))
    c.drawRightString(width - 62, height - 53, normalize_spanish(presentation.presentation_type).upper())
    c.setFont("Helvetica", 8)
    c.setFillColor(colors.HexColor("#64748b"))
    c.drawRightString(width - 62, 32, f"Página {page_number:02d} | Operia Forensic Legal Report")


def _premium_chip(c: canvas.Canvas, x: float, y: float, text: str, tone: str = "info") -> float:
    palette = {
        "info": ("#eef6ff", "#bfdbfe", "#1d4ed8"),
        "valid": ("#ecfdf5", "#bbf7d0", "#047857"),
        "warn": ("#fffbeb", "#fde68a", "#a16207"),
        "risk": ("#fef2f2", "#fecaca", "#b91c1c"),
        "neutral": ("#f8fafc", "#cbd5e1", "#334155"),
    }
    fill, stroke, ink = palette.get(tone, palette["info"])
    text = normalize_spanish(text)
    w = stringWidth(text.upper(), "Helvetica-Bold", 7.5) + 20
    c.setFillColor(colors.HexColor(fill))
    c.setStrokeColor(colors.HexColor(stroke))
    c.roundRect(x, y - 14, w, 22, 7, fill=1, stroke=1)
    c.setFillColor(colors.HexColor(ink))
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(x + 10, y - 7, text.upper())
    return w


def _premium_stat(c: canvas.Canvas, x: float, y: float, w: float, label: str, value: str, tone: str) -> None:
    colors_by_tone = {"valid": "#059669", "info": "#0e7490", "warn": "#b7791f", "risk": "#dc2626", "neutral": "#475569"}
    accent = colors_by_tone.get(tone, "#0e7490")
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor("#d8e4ee"))
    c.roundRect(x, y, w, 66, 13, fill=1, stroke=1)
    c.setFillColor(colors.HexColor(accent))
    c.roundRect(x, y, 6, 66, 3, fill=1, stroke=0)
    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(colors.HexColor("#64748b"))
    c.drawString(x + 18, y + 43, normalize_spanish(label).upper())
    c.setFont("Helvetica-Bold", 16)
    c.setFillColor(colors.HexColor("#0f172a"))
    c.drawString(x + 18, y + 18, normalize_spanish(value))


def _premium_panel(c: canvas.Canvas, x: float, y: float, w: float, h: float, title: str, tone: str = "info") -> None:
    accent = {"info": "#0e7490", "valid": "#059669", "warn": "#b7791f", "risk": "#dc2626", "neutral": "#64748b"}.get(tone, "#0e7490")
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor("#d8e4ee"))
    c.roundRect(x, y, w, h, 14, fill=1, stroke=1)
    c.setFillColor(colors.HexColor("#f8fbff"))
    c.roundRect(x + 1, y + h - 34, w - 2, 33, 13, fill=1, stroke=0)
    c.setFillColor(colors.HexColor(accent))
    c.roundRect(x + 14, y + h - 24, 6, 12, 3, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#334155"))
    c.setFont("Helvetica-Bold", 8)
    c.drawString(x + 28, y + h - 22, normalize_spanish(title).upper())


def _slide_tone(title: str, text: str) -> str:
    haystack = normalize_spanish(f"{title} {text}").lower()
    if any(term in haystack for term in ["riesgo", "limitación", "pendiente", "advertencia"]):
        return "warn"
    if any(term in haystack for term in ["conclusión", "defendible", "validado"]):
        return "valid"
    if any(term in haystack for term in ["daño", "cuantificación", "monto"]):
        return "warn"
    if any(term in haystack for term in ["jurídico", "ley", "norma"]):
        return "info"
    return "neutral"


def _draw_bullets(c: canvas.Canvas, items: list[str], x: float, y: float, width: float, max_items: int, font_size: int = 11) -> float:
    for item in items[:max_items]:
        item = normalize_spanish(item.strip())
        if not item:
            continue
        c.setFillColor(colors.HexColor("#0f766e"))
        c.circle(x, y + 4, 3, fill=1, stroke=0)
        y = _draw_wrapped(c, item, x + 16, y, width - 18, "Helvetica", font_size, font_size + 5, colors.HexColor("#25324a"), max_lines=3) - 5
    return y


def _draw_premium_slide(c: canvas.Canvas, presentation: Presentation, slide, width: float, height: float, page_number: int) -> None:
    _premium_bg(c, width, height)
    _premium_header(c, presentation, width, height, page_number)
    tone = _slide_tone(slide.title, slide.text)

    c.setFont("Helvetica-Bold", 9)
    c.setFillColor(colors.HexColor("#0f766e"))
    c.drawString(68, height - 105, f"{page_number:02d}")
    y = _draw_wrapped(c, normalize_spanish(slide.title), 104, height - 98, width - 300, "Helvetica-Bold", 25, 30, colors.HexColor("#0f172a"), max_lines=2)
    _premium_chip(c, width - 202, height - 101, normalize_spanish(slide.confidentiality_level), "neutral" if slide.confidentiality_level == "interno" else ("risk" if slide.confidentiality_level == "confidencial" else "valid"))

    body = normalize_spanish(slide.text or "Contenido pendiente de completar.")
    raw_items = [line.strip("- • ") for line in body.splitlines() if line.strip()]
    if len(raw_items) <= 1:
        raw_items = [part.strip() for part in re.split(r"(?<=[.;])\s+", body) if part.strip()]

    main_x, main_y, main_w, main_h = 68, 114, width - 316, height - 244
    side_x, side_w = main_x + main_w + 22, 204
    _premium_panel(c, main_x, main_y, main_w, main_h, "Contenido ejecutivo", tone)
    _draw_bullets(c, raw_items, main_x + 30, main_y + main_h - 62, main_w - 58, 8, 12)

    _premium_panel(c, side_x, main_y + main_h - 122, side_w, 122, "Clave de lectura", "info")
    cues = []
    text_l = body.lower()
    if "hash" in text_l or "sha-256" in text_l:
        cues.append("Integridad verificable")
    if "fuente pendiente" in text_l or "pendiente" in text_l:
        cues.append("Requiere verificación")
    if "compatible" in text_l or "limitaciones" in text_l:
        cues.append("Lenguaje defendible")
    if "daño" in text_l or "monto" in text_l:
        cues.append("Impacto económico")
    if not cues:
        cues = ["Síntesis para exposición", "Revisión humana requerida"]
    _draw_bullets(c, cues, side_x + 24, main_y + main_h - 62, side_w - 44, 4, 9)

    _premium_panel(c, side_x, main_y + 78, side_w, 150, "Notas del expositor", "valid")
    note = normalize_spanish(slide.speaker_notes or "Agregar explicación oral: qué demuestra, qué no demuestra y cómo se conecta con la hipótesis.")
    _draw_wrapped(c, note, side_x + 20, main_y + 184, side_w - 40, "Helvetica", 9, 13, colors.HexColor("#14532d"), max_lines=7)

    _premium_panel(c, side_x, main_y, side_w, 58, "Control probatorio", "warn")
    control = "Evitar atribución personal sin evidencia complementaria."
    _draw_wrapped(c, control, side_x + 20, main_y + 26, side_w - 40, "Helvetica", 8.5, 11, colors.HexColor("#713f12"), max_lines=2)


def export_presentation_pdf(presentation: Presentation) -> Path:
    path = new_export_path(f"presentation-{presentation.id}.pdf")
    c = canvas.Canvas(str(path), pagesize=landscape(A4))
    width, height = landscape(A4)
    slides = sorted(presentation.slides, key=lambda item: item.slide_number)
    case = presentation.case
    logo = _logo_path()
    total_damage = sum(item.amount for item in case.damage_items) if case else 0

    _premium_bg(c, width, height)
    if logo:
        c.drawImage(ImageReader(str(logo)), 80, height - 160, width=305, height=112, preserveAspectRatio=True, mask="auto")
    _premium_chip(c, 82, height - 190, "Deck ejecutivo premium", "valid")
    _premium_chip(c, 232, height - 190, "Revisión humana requerida", "warn")
    _premium_chip(c, 410, height - 190, normalize_spanish(presentation.status), "info")

    title = normalize_spanish(case.title if case else presentation.title)
    y = _draw_wrapped(c, title, 82, height - 248, width - 164, "Helvetica-Bold", 34, 39, colors.HexColor("#0f172a"), max_lines=2)
    subtitle = normalize_spanish(f"Caso {case.case_code if case else '-'} | {presentation.presentation_type}")
    _draw_wrapped(c, subtitle, 84, y - 8, width - 168, "Helvetica", 13, 18, colors.HexColor("#475569"), max_lines=2)

    _premium_panel(c, 82, 172, width - 164, 122, "Teoría ejecutiva del caso", "info")
    intro = normalize_spanish(slides[0].text if slides else presentation.title)
    _draw_wrapped(c, intro, 108, 248, width - 216, "Helvetica", 12, 17, colors.HexColor("#25324a"), max_lines=5)

    gap = 16
    stat_w = (width - 164 - gap * 3) / 4
    _premium_stat(c, 82, 86, stat_w, "Evidencias", str(len(case.evidence) if case else 0), "info")
    _premium_stat(c, 82 + stat_w + gap, 86, stat_w, "Hallazgos", str(len(case.findings) if case else 0), "valid")
    _premium_stat(c, 82 + (stat_w + gap) * 2, 86, stat_w, "Daño estimado", f"{total_damage:,.0f} ARS", "warn")
    _premium_stat(c, 82 + (stat_w + gap) * 3, 86, stat_w, "Score deck", str(presentation.review_score or 0), "neutral")
    c.setFont("Helvetica", 8)
    c.setFillColor(colors.HexColor("#64748b"))
    c.drawRightString(width - 82, 56, "Operia Forensic Legal Report | Presentación asistida por IA, no sustituye revisión profesional")
    c.showPage()

    content_slides = slides[1:] if len(slides) > 1 and "portada" in normalize_spanish(slides[0].title).lower() else slides
    for index, slide in enumerate(content_slides, start=1):
        _draw_premium_slide(c, presentation, slide, width, height, index)
        c.showPage()

    _premium_bg(c, width, height)
    _premium_header(c, presentation, width, height, len(content_slides) + 1)
    _premium_chip(c, 82, height - 112, "Cierre profesional", "valid")
    _draw_wrapped(c, "Conclusión para exposición oral", 82, height - 158, width - 164, "Helvetica-Bold", 30, 35, colors.HexColor("#0f172a"), max_lines=2)
    _premium_panel(c, 82, 150, width - 164, 220, "Conclusión defendible", "valid")
    closing = (
        "La presentación sintetiza los elementos centrales del informe, con alcance preliminar, "
        "limitaciones expresas y necesidad de revisión humana profesional. Las afirmaciones fuertes "
        "deben sostenerse con evidencia preservada, cadena de custodia y fuentes verificadas."
    )
    _draw_wrapped(c, closing, 112, 308, width - 224, "Helvetica", 15, 22, colors.HexColor("#25324a"), max_lines=7)
    _draw_wrapped(c, "Próximo paso sugerido: revisar recomendaciones del revisor automático y regenerar el PDF final.", 112, 205, width - 224, "Helvetica-Bold", 12, 18, colors.HexColor("#0f766e"), max_lines=2)
    c.showPage()

    c.save()
    presentation.exported_file = str(path)
    return path



def _export_native_pdf(report: Report) -> Path:
    """Render only the preserved report; never rebuild its facts from a live case."""
    path = new_export_path(f"report-{report.id}.pdf")
    style = ParagraphStyle("NativePreserved", fontName="Helvetica", fontSize=9, leading=12,
                           splitLongWords=True, spaceAfter=6)
    story = []
    for line in report.markdown_content.splitlines():
        story.append(Paragraph(escape(line), style) if line else Spacer(1, 6))
    SimpleDocTemplate(str(path), pagesize=A4, title=report.title, author=report.author).build(story)
    return path
